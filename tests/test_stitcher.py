import os
import re
import struct
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.shared import Inches

from docx_stitcher.cli import main
from docx_stitcher.core import (
    StitchError,
    collect_inputs,
    find_docx_files,
    looks_like_copy,
    next_free_path,
    sort_files,
    stitch,
    trailing_number,
)


def tiny_png(path: Path) -> Path:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    raw = b"\x00\xff\x00\x00"  # one red pixel
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path


def make_doc(path: Path, text: str, image: Path | None = None) -> Path:
    doc = Document()
    doc.add_paragraph(text)
    if image:
        doc.add_picture(str(image), width=Inches(1))
    doc.save(str(path))
    return path


def texts(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs if p.text]


def page_breaks(path: Path) -> int:
    body = Document(str(path)).element.body
    return sum(
        1 for br in body.iter() if br.tag.endswith("}br") and br.get(f"{{{br.nsmap['w']}}}type") == "page"
    )


@pytest.fixture
def chapters(tmp_path: Path) -> Path:
    for name in ["part_10.docx", "part_2.docx", "part_001.docx", "cover.docx"]:
        make_doc(tmp_path / name, Path(name).stem)
    (tmp_path / "~$part_2.docx").write_bytes(b"lock file")
    (tmp_path / "notes.txt").write_text("not a docx")
    return tmp_path


def names(paths):
    return [p.name for p in paths]


def test_find_skips_lock_and_other_files(chapters: Path):
    assert sorted(names(find_docx_files(chapters))) == [
        "cover.docx",
        "part_001.docx",
        "part_10.docx",
        "part_2.docx",
    ]


def test_find_exclude_and_recursive(chapters: Path):
    sub = chapters / "sub"
    sub.mkdir()
    make_doc(sub / "part_3.docx", "nested")
    assert "part_3.docx" not in names(find_docx_files(chapters))
    assert "part_3.docx" in names(find_docx_files(chapters, recursive=True))
    assert "part_10.docx" not in names(find_docx_files(chapters, exclude=["*_10*"]))


def test_sort_by_number_handles_zero_padding_and_missing_numbers(chapters: Path):
    ordered = sort_files(find_docx_files(chapters), "number")
    assert names(ordered) == ["cover.docx", "part_001.docx", "part_2.docx", "part_10.docx"]


def test_sort_number_ignores_prefix():
    files = [Path("intro_3.docx"), Path("body_1.docx"), Path("end_2.docx")]
    assert names(sort_files(files, "number")) == ["body_1.docx", "end_2.docx", "intro_3.docx"]


@pytest.mark.parametrize(
    ("name", "copy", "number"),
    [
        ("part_3 (1).docx", True, 3),
        ("part_3 - Copy.docx", True, 3),
        ("part_3 copy 2.docx", True, 3),
        ("part_3.docx", False, 3),
        ("Report (2021).docx", True, None),
        ("cover.docx", False, None),
    ],
)
def test_copy_detection(name: str, copy: bool, number: int | None):
    assert looks_like_copy(Path(name)) is copy
    assert trailing_number(Path(name)) == number


def test_copies_sort_next_to_original_and_can_be_skipped(chapters: Path):
    make_doc(chapters / "part_2 (1).docx", "dupe")
    ordered = sort_files(find_docx_files(chapters), "number")
    assert names(ordered).index("part_2 (1).docx") == names(ordered).index("part_2.docx") + 1
    assert "part_2 (1).docx" not in names(find_docx_files(chapters, skip_copies=True))


def test_sort_by_name_is_natural():
    files = [Path("Chapter 10.docx"), Path("chapter 9.docx"), Path("Appendix.docx")]
    assert names(sort_files(files, "name")) == ["Appendix.docx", "chapter 9.docx", "Chapter 10.docx"]


def test_sort_modified_and_reverse(tmp_path: Path):
    a, b = make_doc(tmp_path / "a.docx", "a"), make_doc(tmp_path / "b.docx", "b")
    os.utime(a, (2_000_000_000, 2_000_000_000))
    os.utime(b, (1_000_000_000, 1_000_000_000))
    assert names(sort_files([a, b], "modified")) == ["b.docx", "a.docx"]
    assert names(sort_files([a, b], "modified", reverse=True)) == ["a.docx", "b.docx"]


def test_collect_keeps_explicit_file_order_and_skips_output(chapters: Path):
    out = chapters / "merged.docx"
    make_doc(out, "old output")
    files = collect_inputs([chapters / "part_10.docx", chapters, chapters / "cover.docx"], output=out)
    assert names(files) == ["part_10.docx", "cover.docx", "part_001.docx", "part_2.docx"]


def test_collect_rejects_bad_paths(tmp_path: Path):
    with pytest.raises(StitchError, match="No such file"):
        collect_inputs([tmp_path / "missing.docx"])
    (tmp_path / "x.doc").write_text("old format")
    with pytest.raises(StitchError, match="Not a .docx"):
        collect_inputs([tmp_path / "x.doc"])


def test_stitch_merges_in_order_with_page_breaks(chapters: Path):
    files = sort_files(find_docx_files(chapters), "number")
    out = stitch(files, chapters / "merged.docx")
    assert texts(out) == ["cover", "part_001", "part_2", "part_10"]
    assert page_breaks(out) == 3


def test_stitch_without_page_breaks(chapters: Path):
    out = stitch(find_docx_files(chapters), chapters / "merged.docx", page_breaks=False)
    assert page_breaks(out) == 0


def test_stitch_keeps_images(tmp_path: Path):
    png = tiny_png(tmp_path / "dot.png")
    a = make_doc(tmp_path / "a.docx", "first", png)
    b = make_doc(tmp_path / "b.docx", "second", png)
    out = stitch([a, b], tmp_path / "merged.docx")
    doc = Document(str(out))
    assert len(doc.inline_shapes) == 2
    image_parts = [r for r in doc.part.rels.values() if "image" in r.reltype]
    assert image_parts, "images must be carried into the merged package"


def rich_doc(path: Path, n: int, image: Path) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = f"Header {n}"
    doc.add_heading(f"Chapter {n}", 1)
    for style in ["List Number", "List Number", "List Bullet", "Intense Quote"]:
        doc.add_paragraph(f"{style} {n}", style=style)
    doc.add_table(rows=2, cols=2, style="Light Grid Accent 1").cell(0, 0).text = "cell"
    doc.add_picture(str(image), width=Inches(1))
    if n % 2:
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.sections[-1].footer.paragraphs[0].text = f"Footer {n}"
        doc.add_paragraph("second section")
    doc.save(str(path))
    return path


def test_fast_composer_matches_docxcompose(tmp_path: Path):
    from docxcompose.composer import Composer

    from docx_stitcher.core import _fast_composer

    png = tiny_png(tmp_path / "dot.png")
    files = [rich_doc(tmp_path / f"d{n}.docx", n, png) for n in range(5)]

    def merge(factory, out: Path) -> dict[str, bytes]:
        master = Document(str(files[0]))
        composer = factory(master)
        for path in files[1:]:
            master.add_page_break()
            composer.append(Document(str(path)))
        composer.save(str(out))
        with zipfile.ZipFile(out) as z:
            # list definitions get a random nsid on every merge
            return {n: re.sub(rb'w:nsid w:val="\w+"', b"", z.read(n)) for n in z.namelist()}

    assert merge(_fast_composer, tmp_path / "fast.docx") == merge(Composer, tmp_path / "stock.docx")


def test_stitch_without_images(tmp_path: Path):
    png = tiny_png(tmp_path / "dot.png")
    files = [make_doc(tmp_path / f"{n}.docx", f"text {n}", png) for n in range(3)]
    out = stitch(files, tmp_path / "merged.docx", keep_images=False)

    doc = Document(str(out))
    assert len(doc.inline_shapes) == 0
    assert texts(out) == ["text 0", "text 1", "text 2"]
    with zipfile.ZipFile(out) as z:
        assert [n for n in z.namelist() if n.startswith("word/media/")] == []


def test_cli_no_images(tmp_path: Path):
    png = tiny_png(tmp_path / "dot.png")
    for n in range(2):
        make_doc(tmp_path / f"{n}.docx", f"text {n}", png)
    assert main([str(tmp_path), "-q", "--no-images"]) == 0
    assert len(Document(str(tmp_path / "merged.docx")).inline_shapes) == 0


def test_stitch_refuses_to_overwrite(chapters: Path):
    out = make_doc(chapters / "merged.docx", "existing")
    with pytest.raises(StitchError, match="already exists"):
        stitch([chapters / "cover.docx"], out)
    stitch([chapters / "cover.docx"], out, overwrite=True)
    assert texts(out) == ["cover"]


def test_stitch_reports_corrupt_file(tmp_path: Path):
    good = make_doc(tmp_path / "good.docx", "ok")
    bad = tmp_path / "bad.docx"
    bad.write_text("this is not a zip")
    with pytest.raises(StitchError, match="bad.docx"):
        stitch([good, bad], tmp_path / "merged.docx")
    assert not (tmp_path / "merged.docx").exists()


def test_stitch_reports_truncated_file(tmp_path: Path):
    """A file copied only part way is caught before merging, not half way through."""
    good = make_doc(tmp_path / "good.docx", "ok")
    cut = tmp_path / "cut.docx"
    cut.write_bytes(make_doc(tmp_path / "whole.docx", "whole").read_bytes()[:400])
    with pytest.raises(StitchError, match="can't be read"):
        stitch([good, cut], tmp_path / "merged.docx")


def test_stitch_can_skip_unreadable_files(tmp_path: Path):
    files = [make_doc(tmp_path / f"{n}.docx", f"text {n}") for n in range(3)]
    cut = tmp_path / "cut.docx"
    cut.write_bytes(files[0].read_bytes()[:400])
    skipped = []
    out = stitch(
        [files[0], cut, files[1], files[2]],
        tmp_path / "merged.docx",
        skip_unreadable=True,
        on_skip=lambda path, why: skipped.append((path.name, why)),
    )
    assert texts(out) == ["text 0", "text 1", "text 2"]
    assert [name for name, _ in skipped] == ["cut.docx"]


def half_writing_composer(monkeypatch, *, raise_after: bool):
    """Make the merge write a partial file, as it would if it ran out of memory."""
    from docx_stitcher import core

    real = core._fast_composer

    class Dying:
        def __init__(self, master):
            self._composer = real(master)

        def append(self, doc):
            self._composer.append(doc)

        def save(self, filename):
            Path(filename).write_bytes(b"PK\x03\x04 half a document")
            if raise_after:
                raise MemoryError

    monkeypatch.setattr(core, "_fast_composer", Dying)


def test_failed_save_leaves_no_half_written_file(tmp_path: Path, monkeypatch):
    files = [make_doc(tmp_path / f"{n}.docx", f"text {n}") for n in range(2)]
    out = make_doc(tmp_path / "merged.docx", "previous merge")
    half_writing_composer(monkeypatch, raise_after=True)

    with pytest.raises(MemoryError):
        stitch(files, out, overwrite=True)
    assert texts(out) == ["previous merge"], "the earlier file must survive untouched"
    assert list(tmp_path.glob(".*")) == [], "no leftover part file"


def test_incomplete_save_is_reported(tmp_path: Path, monkeypatch):
    files = [make_doc(tmp_path / f"{n}.docx", f"text {n}") for n in range(2)]
    out = tmp_path / "merged.docx"
    half_writing_composer(monkeypatch, raise_after=False)

    with pytest.raises(StitchError, match="not written completely"):
        stitch(files, out)
    assert not out.exists()


def test_progress_callback(chapters: Path):
    seen = []
    stitch(
        find_docx_files(chapters), chapters / "merged.docx", on_progress=lambda d, t, _p: seen.append((d, t))
    )
    assert seen == [(0, 4), (1, 4), (2, 4), (3, 4), (4, 4)]


def test_next_free_path(tmp_path: Path):
    target = tmp_path / "merged.docx"
    assert next_free_path(target) == target
    target.touch()
    (tmp_path / "merged (2).docx").touch()
    assert next_free_path(target).name == "merged (3).docx"


def test_cli_dry_run_writes_nothing(chapters: Path, capsys):
    assert main([str(chapters), "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert out.index("cover.docx") < out.index("part_001.docx") < out.index("part_10.docx")
    assert not (chapters / "merged.docx").exists()


def test_cli_merge_then_rerun(chapters: Path, capsys):
    assert main([str(chapters), "-q"]) == 0
    assert texts(chapters / "merged.docx") == ["cover", "part_001", "part_2", "part_10"]

    assert main([str(chapters), "-q"]) == 1
    assert "already exists" in capsys.readouterr().err

    assert main([str(chapters), "-q", "--rename"]) == 0
    assert texts(chapters / "merged (2).docx") == ["cover", "part_001", "part_2", "part_10"]


def test_cli_files_and_output(chapters: Path):
    out = chapters / "book" / "book.docx"
    assert main([str(chapters / "part_2.docx"), str(chapters / "cover.docx"), "-o", str(out), "-q"]) == 0
    assert texts(out) == ["part_2", "cover"]


def test_cli_no_files(tmp_path: Path, capsys):
    assert main([str(tmp_path)]) == 1
    assert "No .docx files" in capsys.readouterr().err


def test_page_break_helper_counts_manual_breaks(tmp_path: Path):
    doc = Document()
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(str(tmp_path / "b.docx"))
    assert page_breaks(tmp_path / "b.docx") == 1
